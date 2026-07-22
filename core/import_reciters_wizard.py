"""
Logic for the web-based biography import wizard (admin view).
Handles CSV parsing, pre-scan analysis, and the commit transaction.
"""

import csv
import difflib
import io
import re
from collections import defaultdict

from django.db import transaction

from .arabic_date_parser import parse_arabic_date
from .csv_import import _normalize_ar, _find_location_candidates
from .models import Attribute, Biography, Country, Location, Reading, Source, TeacherStudentRelationship
from .search import normalize_arabic

_COLUMN_KEYWORDS = {
    "alias":         ["شهره", "شهرة"],
    "full_name":     ["اسم الكامل", "الاسم الكامل"],
    "birth_date":    ["تاريخ الميلاد", "ميلاد"],
    "birth_city":    ["مدينة الميلاد", "مدينه الميلاد"],
    "birth_country": ["دولة الميلاد", "دوله الميلاد", "بلد الميلاد"],
    "death_date":    ["تاريخ الوفاة", "وفاة"],
    "death_city":    ["مدينة الوفاة", "مدينه الوفاة"],
    "death_country": ["دولة الوفاة", "دوله الوفاة", "بلد الوفاة"],
    "teachers":      ["شيوخ", "شيوخه"],
    "students":      ["تلاميذ", "تلاميذه"],
    "attributes":    ["صفات", "تصنيفات"],
    "source":        ["مصدر"],
    "status":        ["حاله الترجمه", "حالة الترجمة", "معتمد"],
}

_ATTR_SEPS = re.compile(r"[،,;؛/]")
_NAME_SEPS = re.compile(r"[،,;؛\n]")
_DASHES = ("-", "–", "")


def _clean_cell(text):
    text = (text or "").strip()
    return "" if text in _DASHES else text


# Keys are resolved in this order so the specific city/country columns claim
# their header before the greedy bare "ميلاد"/"وفاة" date keywords can match it
# (both "مدينة الميلاد" and "دولة الميلاد" contain "ميلاد").
_DETECTION_ORDER = [
    "alias", "full_name",
    "birth_city", "birth_country", "death_city", "death_country",
    "birth_date", "death_date",
    "teachers", "students", "attributes", "source", "status",
]


def _detect_columns(headers):
    mapping = {}
    used = set()
    norm_headers = [_normalize_ar(h) for h in headers]
    for key in _DETECTION_ORDER:
        keywords = _COLUMN_KEYWORDS[key]
        for i, nh in enumerate(norm_headers):
            if i in used:
                continue
            if any(_normalize_ar(kw) in nh for kw in keywords):
                mapping[key] = i
                used.add(i)
                break
    return mapping


def _loc_from_cols(city, country):
    """Build a location spec from separate city and country cells.

    Returns None when both are empty. Either side may be empty: a country with
    no city, or a city with no (yet) country. A region (e.g. الحجاز، الشام)
    simply arrives in the country column and is treated like any country.
    """
    city = _clean_cell(city)
    country = _clean_cell(country)
    if not city and not country:
        return None
    return {"city_ar": city, "country_ar": country}


def _split_names(text):
    if not text or text.strip() in ("-", "–", ""):
        return []
    return [p.strip() for p in _NAME_SEPS.split(text) if p.strip() and p.strip() not in ("-", "–")]


def _split_attrs(text):
    if not text or text.strip() in ("-", "–", ""):
        return []
    return [p.strip() for p in _ATTR_SEPS.split(text) if p.strip() and p.strip() not in ("-", "–")]


_READING_RE = re.compile(r"^(.*?)\s*\(([^)]+)\)\s*$")


def _parse_name_and_reading(entry):
    """Split 'Name (reading_ar)' into (name, reading_ar). Returns (entry, None) if no bracket."""
    m = _READING_RE.match(entry)
    if m:
        return m.group(1).strip(), m.group(2).strip()
    return entry, None


def _date_is_ambiguous(parsed):
    if parsed["calendar"] is None:
        return False
    if parsed["year"] is None:
        return True
    if parsed["approximate"] and "وقيل" in parsed["raw"]:
        return True
    return False


def _fuzzy_match_attr(name, all_attrs, threshold=0.75):
    norm = _normalize_ar(name)
    best_score, best_attr = 0.0, None
    for attr in all_attrs:
        for candidate in (attr.short_name_ar, attr.long_name_ar, attr.short_name_en):
            if not candidate:
                continue
            nc = _normalize_ar(candidate)
            if nc == norm:
                return 1.0, attr
            score = difflib.SequenceMatcher(None, norm, nc).ratio()
            if score > best_score:
                best_score, best_attr = score, attr
    if best_score >= threshold:
        return best_score, best_attr
    return best_score, None


def _is_approved(status_text):
    normalized = _normalize_ar(status_text or "")
    return "معتمد" in normalized or "approved" in normalized.lower()


# Column headers for the downloadable example template. Each header contains the
# keyword that ``_detect_columns`` matches on, so the filled file uploads back
# cleanly. Order is display order (RTL).
_TEMPLATE_COLUMNS = [
    ("full_name",     "الاسم الكامل"),
    ("alias",         "الشهرة"),
    ("birth_date",    "تاريخ الميلاد"),
    ("birth_city",    "مدينة الميلاد"),
    ("birth_country", "دولة الميلاد"),
    ("death_date",    "تاريخ الوفاة"),
    ("death_city",    "مدينة الوفاة"),
    ("death_country", "دولة الوفاة"),
    ("teachers",      "أبرز شيوخه"),
    ("students",      "أبرز تلاميذه"),
    ("attributes",    "صفات وتصنيفات"),
    ("source",        "المصدر"),
    ("status",        "حالة الترجمة"),
]

# Two illustrative rows: a full city+country entry, then the special cases —
# a region in the country column and a country with no city.
_TEMPLATE_EXAMPLE_ROWS = [
    {
        "full_name":     "عاصم بن أبي النجود الأسدي",
        "alias":         "عاصم الكوفي",
        "birth_date":    "",
        "birth_city":    "الكوفة",
        "birth_country": "العراق",
        "death_date":    "١٢٧هـ",
        "death_city":    "الكوفة",
        "death_country": "العراق",
        "teachers":      "زر بن حبيش، أبو عبد الرحمن السلمي",
        "students":      "حفص بن سليمان، شعبة بن عياش",
        "attributes":    "١٠ك، قراء الكوفة",
        "source":        "غاية النهاية في طبقات القراء",
        "status":        "معتمد",
    },
    {
        "full_name":     "مثال: منطقة بدون مدينة، ووفاة بدولة بدون مدينة",
        "alias":         "",
        "birth_date":    "",
        "birth_city":    "",            # unknown city — region only in the country column
        "birth_country": "الحجاز",      # a region is entered like any country
        "death_date":    "",
        "death_city":    "",            # unknown city
        "death_country": "مصر",         # country with no city
        "teachers":      "",
        "students":      "",
        "attributes":    "",
        "source":        "",
        "status":        "غير معتمد",
    },
]


def build_template_xlsx():
    """Return the bytes of an .xlsx example for the reciters wizard.

    The sheet carries the exact Arabic headers the wizard detects, two example
    rows (covering city+country, region-only and country-only locations), and
    advisory dropdowns (country + approval status) sourced from the database so
    the user can enter data cleanly. The filled file uploads back into the
    wizard directly (the upload step also accepts .xlsx).
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    keys = [key for key, _h in _TEMPLATE_COLUMNS]
    headers = [h for _k, h in _TEMPLATE_COLUMNS]

    wb = Workbook()
    ws = wb.active
    ws.title = "التراجم"
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for example in _TEMPLATE_EXAMPLE_ROWS:
        ws.append([example.get(k, "") for k in keys])

    # Hidden sheet holding dropdown option lists.
    lists = wb.create_sheet("Lists")
    countries = sorted({
        c for c in Country.objects.values_list("name_ar", flat=True) if c and c.strip()
    })
    if not countries:  # fresh DB — fall back to any country text on locations
        countries = sorted({
            c.strip() for c in Location.objects.values_list("country_ar", flat=True)
            if c and c.strip()
        })

    def _write_list(col_idx, values):
        letter = get_column_letter(col_idx)
        for i, value in enumerate(values, start=1):
            lists.cell(row=i, column=col_idx, value=value)
        if not values:
            return None
        return f"Lists!${letter}$1:${letter}${len(values)}"

    country_ref = _write_list(1, countries)
    status_ref = _write_list(2, ["معتمد", "غير معتمد"])
    lists.sheet_state = "hidden"

    last_row = 500 + 1  # header + data rows the dropdowns cover

    def _add_dropdown(key, source_ref):
        if not source_ref or key not in keys:
            return
        letter = get_column_letter(keys.index(key) + 1)
        dv = DataValidation(
            type="list", formula1=source_ref, allow_blank=True,
            showErrorMessage=False,  # advisory only — new values may be typed
        )
        dv.add(f"{letter}2:{letter}{last_row}")
        ws.add_data_validation(dv)

    _add_dropdown("birth_country", country_ref)
    _add_dropdown("death_country", country_ref)
    _add_dropdown("status", status_ref)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def scan_csv_file(file_obj, filename=""):
    """Read an uploaded Django file, parse it (CSV or .xlsx), scan all rows."""
    name = (filename or getattr(file_obj, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        return _scan_raw_rows(_read_xlsx_rows(file_obj))
    text = file_obj.read()
    if isinstance(text, bytes):
        try:
            text = text.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = text.decode("cp1256", errors="replace")
    return scan_csv_text(text)


def _read_xlsx_rows(file_obj):
    """Read an uploaded .xlsx into a list of rows (each a list of trimmed strings)."""
    from openpyxl import load_workbook

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for raw in ws.iter_rows(values_only=True):
        rows.append(["" if c is None else str(c).strip() for c in raw])
    return rows


def scan_csv_text(text):
    """
    Parse CSV text, scan all rows. Returns (rows, stats, csv_text).

    rows: list of JSON-serializable dicts, one per biography.
    stats: summary counts dict.
    csv_text: the original text (stored in session for re-parse if needed).
    """
    return _scan_raw_rows(list(csv.reader(io.StringIO(text))), text)


def _scan_raw_rows(raw_rows, text=""):
    """Scan already-parsed rows (list of lists). Shared by the CSV and xlsx paths."""
    if len(raw_rows) < 2:
        return [], {"total": 0}, text

    headers = raw_rows[0]
    col = _detect_columns(headers)
    all_attrs = list(Attribute.objects.all())

    rows = []
    stats = defaultdict(int)
    attr_match_cache = {}  # name → (score, match_or_None) — shared across all rows

    for raw in raw_rows[1:]:
        def get(key, _raw=raw):
            idx = col.get(key)
            if idx is None or idx >= len(_raw):
                return ""
            return _raw[idx].strip()

        full_name = get("full_name")
        if not full_name:
            continue

        row = {
            "alias_ar":       get("alias"),
            "full_name_ar":   full_name,
            "birth_date_raw": get("birth_date"),
            "birth_city_raw":    get("birth_city"),
            "birth_country_raw": get("birth_country"),
            "death_date_raw": get("death_date"),
            "death_city_raw":    get("death_city"),
            "death_country_raw": get("death_country"),
            "teachers_raw":   get("teachers"),
            "students_raw":   get("students"),
            "attributes_raw": get("attributes"),
            "source_raw":     get("source"),
            "status_raw":     get("status"),
        }

        row["birth_parsed"] = parse_arabic_date(row["birth_date_raw"])
        row["death_parsed"] = parse_arabic_date(row["death_date_raw"])
        row["birth_needs_confirm"] = _date_is_ambiguous(row["birth_parsed"])
        row["death_needs_confirm"] = _date_is_ambiguous(row["death_parsed"])

        if row["birth_needs_confirm"]:
            stats["ambiguous_dates"] += 1
        if row["death_needs_confirm"]:
            stats["ambiguous_dates"] += 1

        birth_loc = _loc_from_cols(row["birth_city_raw"], row["birth_country_raw"])
        death_loc = _loc_from_cols(row["death_city_raw"], row["death_country_raw"])
        row["birth_loc"] = birth_loc
        row["death_loc"] = death_loc
        # Only a city with no country needs resolving now that country is its own
        # column (unlikely, but the wizard still asks the user just in case).
        row["birth_loc_needs_country"] = bool(birth_loc and birth_loc["city_ar"] and not birth_loc["country_ar"])
        row["death_loc_needs_country"] = bool(death_loc and death_loc["city_ar"] and not death_loc["country_ar"])

        if row["birth_loc_needs_country"]:
            stats["missing_countries"] += 1
        if row["death_loc_needs_country"]:
            stats["missing_countries"] += 1

        attr_names = _split_attrs(row["attributes_raw"])
        row["attr_exact_pks"] = []
        row["attr_fuzzy"] = []
        row["attr_missing"] = []
        seen_names = set()
        for name in attr_names:
            if name in seen_names:
                continue
            seen_names.add(name)
            if name not in attr_match_cache:
                attr_match_cache[name] = _fuzzy_match_attr(name, all_attrs)
            score, match = attr_match_cache[name]
            if score == 1.0:
                row["attr_exact_pks"].append(match.pk)
            elif match is not None:
                row["attr_fuzzy"].append({
                    "name": name,
                    "score": round(score, 3),
                    "match_pk": match.pk,
                    "match_name": match.short_name_ar or "",
                })
                stats["fuzzy_attrs"] += 1
            else:
                row["attr_missing"].append(name)
                stats["missing_attrs"] += 1

        rows.append(row)
        stats["total"] += 1

    return rows, dict(stats), text


def get_unique_missing_cities(rows):
    """Return unique city_ar values that need a country, preserving first-seen order."""
    seen = {}
    for row in rows:
        for loc_key in ("birth_loc", "death_loc"):
            if row.get(f"{loc_key}_needs_country") and row.get(loc_key):
                city = row[loc_key]["city_ar"]
                if city and city not in seen:
                    seen[city] = True
    return list(seen.keys())


def get_unique_fuzzy_attrs(rows):
    """Return unique fuzzy-match attribute items across all rows."""
    seen = {}
    for row in rows:
        for item in row.get("attr_fuzzy", []):
            if item["name"] not in seen:
                seen[item["name"]] = item
    return list(seen.values())


def get_unique_missing_attrs(rows):
    """Return unique no-match attribute names across all rows."""
    seen = {}
    for row in rows:
        for name in row.get("attr_missing", []):
            if name not in seen:
                seen[name] = True
    return list(seen.keys())


def _build_bio_lookup():
    """Load all biographies into memory for fast name matching. Returns (name_to_pk, norm_to_pk, norm_list)."""
    all_bios = list(Biography.objects.values_list("id", "full_name_ar", "alias_ar"))
    name_to_pk = {}
    norm_to_pk = {}
    norm_list = []  # [(normalized_name, pk, display_name)]
    for pk, full_name, alias in all_bios:
        if full_name:
            name_to_pk[full_name] = pk
            n = normalize_arabic(full_name)
            if n not in norm_to_pk:
                norm_to_pk[n] = pk
                norm_list.append((n, pk, full_name))
        if alias:
            na = normalize_arabic(alias)
            if na not in norm_to_pk:
                norm_to_pk[na] = pk
                norm_list.append((na, pk, alias))
    return name_to_pk, norm_to_pk, norm_list


def commit_phase3(rows, published_mode, date_decisions, location_decisions, attr_decisions):
    """
    Phase 3: create/update biographies, locations, and attributes.

    date_decisions    : {"0_birth": parsed_dict, "1_death": parsed_dict, ...}
    location_decisions: {"city_ar": {"city_ar": str, "country_ar": str}}
    attr_decisions    : {"attr_name": {"action": "accept"|"create"|"skip",
                                       "match_pk": int|None, "long_name": str}}

    Returns (stats_dict, bio_map) where bio_map is {full_name_ar: pk}.
    """
    stats = defaultdict(int)
    bio_map = {}  # full_name_ar → pk

    # ── Pre-create / resolve attributes ──────────────────────────────────────
    pending_attr_pks = {}  # attr_name → int pk
    for attr_name, decision in attr_decisions.items():
        action = decision.get("action", "skip")
        if action == "accept":
            match_pk = decision.get("match_pk")
            if match_pk:
                pending_attr_pks[attr_name] = int(match_pk)
        elif action == "create":
            long_name = decision.get("long_name") or attr_name
            attr, created = Attribute.objects.get_or_create(
                short_name_ar=attr_name,
                defaults={"long_name_ar": long_name},
            )
            pending_attr_pks[attr_name] = attr.pk
            if created:
                stats["new_attrs"] += 1

    # ── Location resolver ─────────────────────────────────────────────────────
    all_locations = list(Location.objects.all())
    loc_cache = {}  # (city_ar, country_ar) → Location

    def resolve_location(city_ar, country_ar):
        key = (city_ar or "", country_ar or "")
        if key in loc_cache:
            return loc_cache[key]
        spec = {"city_ar": key[0], "country_ar": key[1]}
        candidates = _find_location_candidates(spec, all_locations)
        if candidates and candidates[0][1]:  # exact normalized match
            loc = candidates[0][2]
        else:
            loc, created = Location.objects.get_or_create(
                city_ar=key[0], country_ar=key[1]
            )
            if created:
                stats["new_locations"] += 1
                all_locations.append(loc)
        loc_cache[key] = loc
        return loc

    # ── Commit biographies ────────────────────────────────────────────────────
    from .merge_utils import find_similar_to

    with transaction.atomic():
        for idx, row in enumerate(rows):
            name = row["full_name_ar"]

            birth_parsed = date_decisions.get(f"{idx}_birth", row["birth_parsed"])
            death_parsed = date_decisions.get(f"{idx}_death", row["death_parsed"])

            existing = None
            similar = find_similar_to(name)
            if similar:
                top_stub, match_type, _, similarity = similar[0]
                if match_type == "cross_field" and similarity == 1.0:
                    try:
                        existing = Biography.objects.get(pk=top_stub.pk)
                    except Biography.DoesNotExist:
                        pass

            bio = existing or Biography()
            is_new = existing is None

            bio.full_name_ar = name
            if row["alias_ar"]:
                bio.alias_ar = row["alias_ar"]

            if published_mode == "published":
                bio.published = True
            elif published_mode == "unpublished":
                if is_new:
                    bio.published = False
            else:
                bio.published = _is_approved(row["status_raw"])

            for prefix, parsed in [("birth", birth_parsed), ("death", death_parsed)]:
                setattr(bio, f"{prefix}_date_raw_ar", parsed["raw"])
                if parsed["calendar"] is None or parsed["year"] is None:
                    continue
                cal = "hijri" if parsed["calendar"] == "hijri" else "greg"
                year_field = f"{prefix}_{cal}_year"
                if is_new or not getattr(bio, year_field):
                    setattr(bio, year_field, parsed["year"])
                    setattr(bio, f"{prefix}_{cal}_month", parsed["month"])
                    setattr(bio, f"{prefix}_{cal}_day", parsed["day"])
                    setattr(bio, f"{prefix}_{cal}_approximate", parsed["approximate"])

            for loc_key, field_name in [("birth_loc", "birthplace"), ("death_loc", "death_location")]:
                raw_loc = row.get(loc_key)
                if not raw_loc:
                    continue
                city_ar = raw_loc["city_ar"]
                country_ar = raw_loc["country_ar"]
                if row.get(f"{loc_key}_needs_country") and city_ar in location_decisions:
                    decided = location_decisions[city_ar]
                    city_ar = decided.get("city_ar", city_ar)
                    country_ar = decided.get("country_ar", "")
                if city_ar or country_ar:
                    loc_obj = resolve_location(city_ar, country_ar)
                    if loc_obj and (is_new or not getattr(bio, field_name + "_id")):
                        setattr(bio, field_name, loc_obj)

            bio.save()

            if row["source_raw"] and row["source_raw"] not in ("-", "–"):
                raw_sources = _NAME_SEPS.split(row["source_raw"])
                existing_names = set(bio.sources.values_list("name_ar", flat=True))
                for src in raw_sources:
                    src = src.strip()
                    if src and src not in existing_names:
                        Source.objects.create(biography=bio, name_ar=src)
                        existing_names.add(src)
                        stats["new_sources"] += 1

            attrs_to_add = list(row.get("attr_exact_pks", []))
            all_req_names = (
                [item["name"] for item in row.get("attr_fuzzy", [])]
                + row.get("attr_missing", [])
            )
            for attr_name in all_req_names:
                pk = pending_attr_pks.get(attr_name)
                if pk:
                    attrs_to_add.append(pk)
            if attrs_to_add:
                bio.attributes.add(*attrs_to_add)

            bio_map[bio.full_name_ar] = bio.pk
            stats["created" if is_new else "updated"] += 1

    return dict(stats), bio_map


def pre_scan_phase4(rows, bio_map):
    """
    Pre-scan Phase 4: find fuzzy matches and unmatched teacher/student names.

    Returns {"fuzzy_matches": [...], "unmatched": [...], "detected_readings": [...]}.
    Each fuzzy item: {person_name, reading_ar, matched_name, matched_pk, score (%), occurrences}
    Each unmatched item: {person_name, reading_ar, occurrences}
    Each detected_readings item: {person_name, reading_ar} for all entries with a parenthetical
    """
    name_to_pk, norm_to_pk, norm_list = _build_bio_lookup()

    fuzzy_seen = {}    # person_name → item dict
    unmatched_seen = {}  # person_name → item dict
    detected_readings = []  # {person_name, reading_ar} for ALL entries where a reading was found

    for row in rows:
        if row["full_name_ar"] not in bio_map:
            continue
        for raw_key in ("teachers_raw", "students_raw"):
            for raw_entry in _split_names(row.get(raw_key, "")):
                person_name, reading_ar = _parse_name_and_reading(raw_entry)
                if not person_name:
                    continue
                if reading_ar:
                    detected_readings.append({"person_name": person_name, "reading_ar": reading_ar})
                # Exact match → no decision needed
                if person_name in name_to_pk:
                    continue
                norm = normalize_arabic(person_name)
                if norm in norm_to_pk:
                    continue
                # Fuzzy search
                best_score, best_pk, best_name = 0.0, None, None
                for bio_norm, bio_pk, bio_name in norm_list:
                    if not bio_norm:
                        continue
                    score = difflib.SequenceMatcher(None, norm, bio_norm).ratio()
                    if score > best_score:
                        best_score, best_pk, best_name = score, bio_pk, bio_name

                if best_score >= 0.85:
                    if person_name in fuzzy_seen:
                        fuzzy_seen[person_name]["occurrences"] += 1
                    else:
                        fuzzy_seen[person_name] = {
                            "person_name": person_name,
                            "reading_ar": reading_ar,
                            "matched_name": best_name,
                            "matched_pk": best_pk,
                            "score": round(best_score * 100, 1),
                            "occurrences": 1,
                        }
                else:
                    if person_name in unmatched_seen:
                        unmatched_seen[person_name]["occurrences"] += 1
                    else:
                        unmatched_seen[person_name] = {
                            "person_name": person_name,
                            "reading_ar": reading_ar,
                            "occurrences": 1,
                        }

    return {
        "fuzzy_matches": list(fuzzy_seen.values()),
        "unmatched": list(unmatched_seen.values()),
        "detected_readings": detected_readings,
    }


def commit_phase4(rows, bio_map, rel_decisions):
    """
    Phase 4: teacher/student relationship linking based on user decisions.

    rel_decisions: {
        "fuzzy": {person_name: {"action": "link"|"stub"|"skip", "matched_pk": int}},
        "unmatched": {person_name: {"action": "stub"|"skip"}},
    }

    Returns stats dict.
    """
    stats = defaultdict(int)

    name_to_pk, norm_to_pk, norm_list_triples = _build_bio_lookup()
    norm_list = [(n, pk) for n, pk, _ in norm_list_triples]

    fuzzy_decisions = rel_decisions.get("fuzzy", {})
    unmatched_decisions = rel_decisions.get("unmatched", {})

    def resolve_person(person_name):
        """Returns (pk_or_None, should_create_stub)."""
        if person_name in name_to_pk:
            return name_to_pk[person_name], False
        norm = normalize_arabic(person_name)
        if norm in norm_to_pk:
            return norm_to_pk[norm], False

        fd = fuzzy_decisions.get(person_name)
        if fd:
            action = fd.get("action", "skip")
            if action == "link":
                mpk = fd.get("matched_pk")
                if mpk:
                    return int(mpk), False
                return None, True
            if action == "stub":
                return None, True
            return None, False  # skip

        ud = unmatched_decisions.get(person_name)
        if ud:
            if ud.get("action") == "skip":
                return None, False
        return None, True  # default: create stub

    for row in rows:
        if row["full_name_ar"] not in bio_map:
            continue
        student_pk = bio_map[row["full_name_ar"]]

        for role, raw_key in [("teacher", "teachers_raw"), ("student", "students_raw")]:
            for raw_entry in _split_names(row.get(raw_key, "")):
                if not raw_entry:
                    continue
                person_name, reading_ar = _parse_name_and_reading(raw_entry)
                if not person_name:
                    continue

                matched_pk, should_create_stub = resolve_person(person_name)

                if matched_pk is None and should_create_stub:
                    stub, created = Biography.objects.get_or_create(
                        full_name_ar=person_name,
                        defaults={"published": False},
                    )
                    if created:
                        stats["stub_bios"] += 1
                        name_to_pk[person_name] = stub.pk
                        n = normalize_arabic(person_name)
                        if n not in norm_to_pk:
                            norm_to_pk[n] = stub.pk
                            norm_list.append((n, stub.pk))
                    matched_pk = stub.pk

                if matched_pk is None:
                    continue

                teacher_pk = matched_pk if role == "teacher" else student_pk
                actual_student_pk = student_pk if role == "teacher" else matched_pk
                if teacher_pk != actual_student_pk:
                    rel, created = TeacherStudentRelationship.objects.get_or_create(
                        teacher_id=teacher_pk,
                        student_id=actual_student_pk,
                    )
                    if reading_ar:
                        reading_obj, _ = Reading.objects.get_or_create(description_ar=reading_ar)
                        rel.notes.add(reading_obj)
                        stats["readings_linked"] += 1
                    if created:
                        stats["relationships"] += 1

    return dict(stats)

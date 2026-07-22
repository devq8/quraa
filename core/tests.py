import csv
import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from . import csv_import
from .models import Attribute, Biography, City, Country, Location


class NormalizedLocationTests(TestCase):
    """The ``Location`` wrapper keeps its FKs and text columns in sync in both
    directions, and a region is stored as an ordinary ``Country``."""

    def test_legacy_text_creation_backfills_fks(self):
        loc = Location.objects.create(
            city_ar="الكوفة", city_en="Kufa",
            country_ar="العراق", country_en="Iraq",
        )
        self.assertIsNotNone(loc.country_id)
        self.assertIsNotNone(loc.city_id)
        self.assertEqual(loc.country.name_ar, "العراق")
        self.assertEqual(loc.city.name_ar, "الكوفة")
        self.assertEqual(loc.city.country_id, loc.country_id)

    def test_region_is_stored_as_country(self):
        makkah = Location.objects.create(
            city_ar="مكة", city_en="Makkah",
            country_ar="الحجاز", country_en="Hijaz",
        )
        madinah = Location.objects.create(
            city_ar="المدينة", city_en="Madinah",
            country_ar="الحجاز", country_en="Hijaz",
        )
        # Both cities point at the single الحجاز Country row.
        self.assertEqual(Country.objects.filter(name_ar="الحجاز").count(), 1)
        self.assertEqual(makkah.city.country, madinah.city.country)
        self.assertEqual(makkah.city.country.name_ar, "الحجاز")

    def test_country_only_location_has_blank_city(self):
        egypt = Country.objects.create(name_ar="مصر", name_en="Egypt")
        loc = Location.objects.create(country=egypt)
        self.assertEqual(loc.city_ar, "")
        self.assertIsNone(loc.city_id)
        self.assertEqual(loc.country_ar, "مصر")

    def test_fk_creation_mirrors_text_columns(self):
        iraq = Country.objects.create(name_ar="العراق", name_en="Iraq")
        basra = City.objects.create(name_ar="البصرة", name_en="Basra", country=iraq)
        loc = Location.objects.create(city=basra)
        # country inferred from the city, text columns mirrored from the FKs.
        self.assertEqual(loc.country_id, iraq.id)
        self.assertEqual(loc.city_ar, "البصرة")
        self.assertEqual(loc.country_ar, "العراق")
        self.assertEqual(loc.city_en, "Basra")


class RecitersTemplateTests(TestCase):
    """The downloadable Excel example round-trips back through the wizard."""

    def test_template_builds_and_scans_back(self):
        import io

        from . import import_reciters_wizard as wiz

        data = wiz.build_template_xlsx()
        self.assertTrue(data)

        rows, stats, _ = wiz.scan_csv_file(io.BytesIO(data), "reciters_import_example.xlsx")
        self.assertEqual(len(rows), 2)
        # Row 1: full city + country.
        self.assertEqual(rows[0]["birth_loc"], {"city_ar": "الكوفة", "country_ar": "العراق"})
        # Row 2: region-only birthplace, country-only death location.
        self.assertEqual(rows[1]["birth_loc"], {"city_ar": "", "country_ar": "الحجاز"})
        self.assertEqual(rows[1]["death_loc"], {"city_ar": "", "country_ar": "مصر"})


class BiographyLiveDuplicateAdminTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin_user = User.objects.create_superuser(
            email="admin-live@example.com",
            password="password",
            username="admin-live@example.com",
        )
        self.url = reverse("admin:core_biography_live_duplicates")
        self.biography = Biography.objects.create(
            full_name_ar="أحمد بن علي الشاطبي",
            full_name_en="Ahmad ibn Ali al-Shatibi",
            alias_ar="الشاطبي",
            alias_en="Al-Shatibi",
        )

    def test_duplicate_results_include_admin_change_url(self):
        self.client.force_login(self.admin_user)

        response = self.client.get(self.url, {"q": "احمد بن علي الشاطبي"})

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data["results"]), 1)
        self.assertEqual(data["results"][0]["id"], self.biography.pk)
        self.assertEqual(
            data["results"][0]["url"],
            reverse("admin:core_biography_change", args=[self.biography.pk]),
        )

    def test_edit_mode_excludes_current_biography(self):
        self.client.force_login(self.admin_user)

        response = self.client.get(
            self.url,
            {"q": "احمد بن علي الشاطبي", "exclude_id": self.biography.pk},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["results"], [])

    def test_alias_input_returns_matches(self):
        self.client.force_login(self.admin_user)

        response = self.client.get(self.url, {"q": "Al-Shatibi", "field": "alias_en"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["results"][0]["id"], self.biography.pk)

    def test_staff_without_view_permission_cannot_access_endpoint(self):
        User = get_user_model()
        staff_user = User.objects.create_user(
            email="staff-live@example.com",
            password="password",
            username="staff-live@example.com",
            is_staff=True,
        )
        self.client.force_login(staff_user)

        response = self.client.get(self.url, {"q": "احمد بن علي الشاطبي"})

        self.assertNotEqual(response.status_code, 200)

    def test_empty_and_short_queries_return_no_matches(self):
        self.client.force_login(self.admin_user)

        empty_response = self.client.get(self.url, {"q": ""})
        short_response = self.client.get(self.url, {"q": "ab"})

        self.assertEqual(empty_response.status_code, 200)
        self.assertEqual(short_response.status_code, 200)
        self.assertEqual(empty_response.json()["results"], [])
        self.assertEqual(short_response.json()["results"], [])


class BiographyExportAdminTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin_user = User.objects.create_superuser(
            email="admin@example.com",
            password="password",
            username="admin@example.com",
        )
        self.client.force_login(self.admin_user)

        self.kufa = Location.objects.create(
            city_ar="الكوفة",
            city_en="Kufa",
            country_ar="العراق",
            country_en="Iraq",
        )
        self.makkah = Location.objects.create(
            city_ar="مكة",
            city_en="Makkah",
            country_ar="الحجاز",
            country_en="Hijaz",
        )
        self.madinah = Location.objects.create(
            city_ar="المدينة",
            city_en="Madinah",
            country_ar="الحجاز",
            country_en="Hijaz",
        )
        self.attribute_10k = Attribute.objects.create(
            short_name_ar="١٠ك",
            short_name_en="10M",
            long_name_ar="القراءات العشر الكبرى",
            long_name_en="Ten Major Readings",
        )
        self.attribute_azhar = Attribute.objects.create(
            short_name_ar="أزهر",
            short_name_en="Azhar",
            long_name_ar="قراء الأزهر",
            long_name_en="Azhar Reciters",
        )

        self.published_biography = Biography.objects.create(
            full_name_ar="عاصم بن أبي النجود",
            full_name_en="Asim ibn Abi al-Najud",
            alias_ar="عاصم الكوفي",
            alias_en="Asim al-Kufi",
            birthplace=self.kufa,
            death_location=self.madinah,
            birth_hijri_year=60,
            birth_hijri_month=3,
            birth_hijri_day=15,
            birth_greg_year=680,
            birth_greg_month=1,
            birth_greg_day=5,
            death_hijri_year=127,
            death_hijri_month=10,
            death_hijri_day=20,
            death_greg_year=745,
            death_greg_month=8,
            death_greg_day=1,
            comments_ar="إمام مشهور",
            comments_en="A famous imam",
            published=True,
        )
        self.published_biography.hometown.set([self.kufa])
        self.published_biography.attributes.set([
            self.attribute_10k,
            self.attribute_azhar,
        ])

        self.draft_biography = Biography.objects.create(
            full_name_ar="قارئ غير منشور",
            full_name_en="Unpublished Reciter",
            birthplace=self.makkah,
            published=False,
        )

    def test_csv_export_uses_import_headers_and_utf8_bom(self):
        response = self.client.get(reverse("admin:core_biography_export_csv"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content[:3], b"\xef\xbb\xbf")
        self.assertIn(
            'attachment; filename="biographies_export_',
            response["Content-Disposition"],
        )

        text = response.content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)

        self.assertEqual(reader.fieldnames, csv_import.sample_columns())
        self.assertEqual(len(rows), 2)

        row = next(r for r in rows if r["full_name_ar"] == "عاصم بن أبي النجود")
        self.assertEqual(row["published"], "yes")
        self.assertEqual(row["birthplace_city_ar"], "الكوفة")
        self.assertEqual(row["hometown_country_en"], "Iraq")
        self.assertEqual(row["death_location_city_en"], "Madinah")
        self.assertEqual(row["attribute_1"], "١٠ك")
        self.assertEqual(row["attribute_2"], "أزهر")
        self.assertEqual(row["comments_ar"], "إمام مشهور")

    def test_xlsx_export_uses_import_headers_and_values(self):
        response = self.client.get(reverse("admin:core_biography_export_xlsx"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        worksheet = workbook["Biographies"]
        rows = list(worksheet.iter_rows(values_only=True))

        self.assertEqual(list(rows[0]), csv_import.sample_columns())
        values = dict(zip(rows[0], rows[1]))
        self.assertEqual(values["full_name_ar"], "عاصم بن أبي النجود")
        self.assertEqual(values["published"], "yes")
        self.assertEqual(values["birthplace_city_en"], "Kufa")
        self.assertEqual(values["hometown_city_ar"], "الكوفة")
        self.assertEqual(values["attribute_1"], "١٠ك")

    def test_csv_export_preserves_admin_filters(self):
        response = self.client.get(
            reverse("admin:core_biography_export_csv"),
            {"published__exact": "1"},
        )

        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["full_name_ar"], "عاصم بن أبي النجود")

    def test_changelist_export_links_preserve_querystring(self):
        response = self.client.get(
            reverse("admin:core_biography_changelist"),
            {"published__exact": "1", "q": "عاصم"},
        )

        self.assertContains(
            response,
            f'{reverse("admin:core_biography_export_csv")}?published__exact=1&amp;q=%D8%B9%D8%A7%D8%B5%D9%85',
        )
        self.assertContains(
            response,
            f'{reverse("admin:core_biography_export_xlsx")}?published__exact=1&amp;q=%D8%B9%D8%A7%D8%B5%D9%85',
        )
